"""Generate the fictional Lantern Lane Goods document set.

Run from the repo root:  python3 scripts/generate_documents.py
Writes 32 files to documents/. Everything is fictional.

Dependencies: python-docx, reportlab, Pillow. The .xlsx files are written with
the standard library only (zipfile + minimal SpreadsheetML).
Python 3.9 compatible.
"""
import os
import random
import zipfile
from xml.sax.saxutils import escape

from docx import Document
from PIL import Image, ImageDraw, ImageFilter, ImageFont
from reportlab.lib.pagesizes import LETTER
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, "documents")
COMPANY = "Lantern Lane Goods"


# ---------- writers ----------

def write_docx(name, title, blocks):
    doc = Document()
    doc.add_heading(title, level=1)
    for b in blocks:
        if isinstance(b, tuple) and b[0] == "h":
            doc.add_heading(b[1], level=2)
        elif isinstance(b, list):
            for item in b:
                doc.add_paragraph(item, style="List Bullet")
        else:
            doc.add_paragraph(b)
    doc.save(os.path.join(OUT, name))


def write_pdf(name, title, blocks):
    styles = getSampleStyleSheet()
    story = [Paragraph(escape(title), styles["Title"]), Spacer(1, 8)]
    for b in blocks:
        if isinstance(b, tuple) and b[0] == "h":
            story.append(Paragraph(escape(b[1]), styles["Heading2"]))
        elif isinstance(b, list):
            for item in b:
                story.append(Paragraph("&bull; " + escape(item), styles["BodyText"]))
        else:
            story.append(Paragraph(escape(b), styles["BodyText"]))
        story.append(Spacer(1, 4))
    SimpleDocTemplate(os.path.join(OUT, name), pagesize=LETTER).build(story)


def col_letter(i):
    s = ""
    i += 1
    while i:
        i, r = divmod(i - 1, 26)
        s = chr(65 + r) + s
    return s


def write_xlsx(name, sheets):
    """sheets: list of (sheet_name, rows). Numbers stay numeric, everything else is an inline string."""
    def cell(ref, v):
        if isinstance(v, (int, float)) and not isinstance(v, bool):
            return '<c r="%s"><v>%s</v></c>' % (ref, v)
        return '<c r="%s" t="inlineStr"><is><t>%s</t></is></c>' % (ref, escape(str(v)))

    ct = ['<?xml version="1.0" encoding="UTF-8" standalone="yes"?>',
          '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">',
          '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>',
          '<Default Extension="xml" ContentType="application/xml"/>',
          '<Override PartName="/xl/workbook.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/>']
    wb_sheets, wb_rels = [], []
    with zipfile.ZipFile(os.path.join(OUT, name), "w", zipfile.ZIP_DEFLATED) as z:
        for i, (sname, rows) in enumerate(sheets, 1):
            xml_rows = []
            for r, row in enumerate(rows, 1):
                cells = "".join(cell("%s%d" % (col_letter(c), r), v) for c, v in enumerate(row) if v is not None)
                xml_rows.append('<row r="%d">%s</row>' % (r, cells))
            z.writestr("xl/worksheets/sheet%d.xml" % i,
                       '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                       '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
                       '<sheetData>%s</sheetData></worksheet>' % "".join(xml_rows))
            ct.append('<Override PartName="/xl/worksheets/sheet%d.xml" '
                      'ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/>' % i)
            wb_sheets.append('<sheet name="%s" sheetId="%d" r:id="rId%d"/>' % (escape(sname), i, i))
            wb_rels.append('<Relationship Id="rId%d" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" '
                           'Target="worksheets/sheet%d.xml"/>' % (i, i))
        ct.append("</Types>")
        z.writestr("[Content_Types].xml", "".join(ct))
        z.writestr("_rels/.rels",
                   '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                   '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
                   '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="xl/workbook.xml"/>'
                   '</Relationships>')
        z.writestr("xl/workbook.xml",
                   '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                   '<workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" '
                   'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
                   '<sheets>%s</sheets></workbook>' % "".join(wb_sheets))
        z.writestr("xl/_rels/workbook.xml.rels",
                   '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                   '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">%s</Relationships>'
                   % "".join(wb_rels))


def find_font(size):
    for p in ("/System/Library/Fonts/Supplemental/Times New Roman.ttf",
              "/System/Library/Fonts/Supplemental/Courier New.ttf",
              "/System/Library/Fonts/Supplemental/Arial.ttf"):
        if os.path.exists(p):
            return ImageFont.truetype(p, size)
    return ImageFont.load_default()


def write_scanned_pdf(name, pages):
    """Image-only PDF: text rendered to a slightly rotated, noisy grayscale image. No text layer."""
    random.seed(7)
    font, bold = find_font(30), find_font(38)
    images = []
    for lines in pages:
        img = Image.new("L", (1700, 2200), 246)
        d = ImageDraw.Draw(img)
        y = 140
        for ln in lines:
            f = bold if ln.startswith("#") else font
            d.text((150, y), ln.lstrip("# "), fill=35, font=f)
            y += 58 if ln else 30
        for _ in range(4000):
            d.point((random.randrange(1700), random.randrange(2200)), fill=random.randrange(120, 200))
        img = img.rotate(0.8, fillcolor=240).filter(ImageFilter.GaussianBlur(0.7))
        images.append(img.convert("RGB"))
    images[0].save(os.path.join(OUT, name), save_all=True, append_images=images[1:], resolution=200)


# ---------- content ----------

def build():
    os.makedirs(OUT, exist_ok=True)
    for f in os.listdir(OUT):
        os.remove(os.path.join(OUT, f))

    # D01 old return policy
    write_docx("Return Policy v2.docx", "%s - Return Policy (v2)" % COMPANY, [
        "Effective: 2024-02-01",
        "We want you to love your home. If something is not right, you can return it.",
        ("h", "Time limit"),
        "Items can be returned within 30 days of delivery. Items must be unused and in original packaging.",
        ("h", "Return shipping"),
        "The customer pays return shipping. We refund the item price once the warehouse receives the return.",
        ("h", "Not returnable"),
        ["Custom or made-to-order furniture", "Gift cards", "Clearance items marked final sale"],
    ])

    # D02 superseded v3
    write_docx("Return Policy v3 FINAL.docx", "%s - Return Policy (v3)" % COMPANY, [
        "Effective: 2026-03-01",
        ("h", "Time limit"),
        "Unused items can be returned within 45 days of delivery, in original packaging.",
        ("h", "Return shipping"),
        "We email a free prepaid return label for orders over $75. For smaller orders, the customer pays return shipping.",
        ("h", "Not returnable"),
        ["Custom or made-to-order furniture", "Gift cards", "Clearance items marked final sale"],
    ])

    # D03 current v3
    write_docx("Return Policy v3 final2.docx", "%s - Return Policy (v3)" % COMPANY, [
        "Effective: 2026-04-15",
        ("h", "Time limit"),
        "Unused items can be returned within 45 days of delivery, in original packaging.",
        ("h", "Return shipping"),
        "We email a free prepaid return label for US orders over $100. For smaller orders, the customer pays return shipping.",
        "International customers (Canada) pay their own return shipping. The free label does not apply outside the US.",
        ("h", "Refunds"),
        "Refunds go back to the original payment method within 5 business days after the warehouse checks the item.",
        ("h", "Not returnable"),
        ["Custom or made-to-order furniture", "Gift cards", "Clearance items marked final sale"],
    ])

    # D04 / D05 shipping rates
    write_xlsx("Shipping Rates 2025.xlsx", [("Rates 2025", [
        ["Destination", "Service", "Rate (USD)", "Free shipping threshold (USD)", "Notes"],
        ["Contiguous US", "Standard", 6.95, 50, "3-7 business days"],
        ["Contiguous US", "Express", 18.95, None, "1-2 business days"],
        ["Alaska / Hawaii", "Standard", 16.95, None, "No free shipping"],
        ["Canada", "Standard", 22.95, None, "Duties paid by customer"],
    ])])
    write_xlsx("Shipping Rates 2026.xlsx", [("Rates 2026", [
        ["Destination", "Service", "Rate (USD)", "Free shipping threshold (USD)", "Notes"],
        ["Contiguous US", "Standard", 7.95, 60, "3-7 business days"],
        ["Contiguous US", "Express", 19.95, None, "1-2 business days"],
        ["Alaska / Hawaii", "Standard", 19.95, None, "Flat rate. No free shipping"],
        ["Canada", "Standard", 24.95, None, "Duties paid by customer"],
    ]), ("Change log", [
        ["Date", "Change", "By"],
        ["2026-01-05", "Standard rate 6.95 -> 7.95, free threshold 50 -> 60", "Ops"],
        ["2026-01-05", "Alaska/Hawaii 16.95 -> 19.95", "Ops"],
    ])])

    # D06
    write_docx("International Shipping FAQ.docx", "International Shipping FAQ", [
        ("h", "Where do you ship?"),
        "We ship within the United States and to Canada. We do not ship to other countries at this time.",
        ("h", "How long does shipping to Canada take?"),
        "Usually 7-14 business days after the order ships.",
        ("h", "Who pays duties and taxes?"),
        "The customer pays any duties, taxes and brokerage fees on delivery.",
        ("h", "Can I return an international order?"),
        "Yes, under our standard return policy. International customers pay return shipping.",
    ])

    # D07
    write_pdf("Warranty Policy.pdf", "%s - Limited Warranty" % COMPANY, [
        "This warranty covers defects in materials and workmanship under normal home use.",
        ("h", "Coverage periods"),
        ["Furniture: 1 year from delivery", "Textiles (bedding, throws, curtains, rugs): 90 days from delivery",
         "Home decor: not covered by warranty (standard returns only)"],
        ("h", "Not covered"),
        ["Normal wear, fading from sunlight, pet damage", "Damage from not following the Product Care Guide",
         "Commercial or rental use"],
        ("h", "How to claim"),
        "Email support with the order number and photos. We repair, replace or refund at our choice.",
    ])

    # D08
    write_docx("Price Match Policy.docx", "Price Match Policy", [
        "If you find the same item for less, we will match the price.",
        ["The request must be made within 14 days of purchase.",
         "The lower price must be from an authorized retailer. Marketplace sellers (for example third-party sellers on large marketplaces) are not eligible.",
         "The item must be identical: same brand, model, size and color, and in stock.",
         "Price match cannot be combined with other promotions, discount codes or sale events."],
        "CS agents: issue the difference as a refund to the original payment method.",
    ])

    # D09
    write_docx("Holiday Returns Memo 2025.docx", "Memo: Holiday return window 2025", [
        "To: All staff    From: Owners    Date: 2025-10-20",
        "For the 2025 holiday season, orders placed from November 1 to December 24, 2025 can be returned until January 31, 2026.",
        "All other return rules stay the same. Please update the CS macros for the season.",
    ])

    # D10 / D11 handbooks
    common_hb = [
        ("h", "Working hours"),
        "Core hours are 10:00-16:00 local time. Tell your manager if you will be offline during core hours.",
        ("h", "Code of conduct"),
        "Treat customers, suppliers and coworkers with respect. Report concerns to an owner.",
    ]
    write_pdf("Employee Handbook 2026.pdf", "%s - Employee Handbook 2026" % COMPANY, [
        "Effective: 2026-01-01. This handbook replaces all earlier versions.",
        ("h", "Paid time off"),
        "Full-time staff get 15 days of paid time off (PTO) per calendar year, plus public holidays.",
        ("h", "Remote work"),
        "Fridays are remote work days for all roles except warehouse staff.",
    ] + common_hb)
    write_pdf("Employee Handbook 2023.pdf", "%s - Employee Handbook 2023" % COMPANY, [
        "Effective: 2023-01-01.",
        ("h", "Paid time off"),
        "Full-time staff get 12 days of paid time off (PTO) per calendar year, plus public holidays.",
        ("h", "Remote work"),
        "Remote work needs manager approval each time.",
    ] + common_hb)

    # D12
    write_docx("Expense Policy.docx", "Expense Policy", [
        ["Travel meals: up to $40 per day.", "Any single expense over $500 needs manager approval before purchase.",
         "Submit receipts within 30 days to the bookkeeper."],
    ])

    # D13 / D32 damaged items (duplicate)
    damaged = ("CS SOP - Damaged Items", [
        "Use this SOP when a customer reports an item that arrived damaged.",
        ["The customer must send photos of the damage and the box within 7 days of delivery.",
         "Offer a replacement or a full refund. Customer chooses.",
         "If the item value is under $40, the customer does not need to send it back.",
         "If the item value is $40 or more, email a prepaid return label. The item must come back before the refund.",
         "Log the case in the damage tracker with the carrier name."],
        "Refund approval limits are in the Refund Escalation SOP.",
    ])
    write_docx("CS SOP - Damaged Items.docx", damaged[0], damaged[1])
    write_docx("damaged items process copy.docx", damaged[0], damaged[1])

    # D14
    write_docx("CS SOP - Refund Escalation.docx", "CS SOP - Refund Escalation", [
        ["CS agents can approve refunds up to $250.", "Refunds over $250 need approval from the CS lead.",
         "Refunds over $1,000 also need an owner to sign off.", "Note the approver's name in the order notes."],
    ])

    # D15 outdated macros
    write_docx("CS Macros (copy-paste replies).docx", "CS Macros", [
        ("h", "Return request"),
        "Hi {name}, thanks for reaching out! You can return unused items within 30 days of delivery. Reply to this email and we will send the steps.",
        ("h", "Where is my order"),
        "Hi {name}, your order ships within 1-2 business days. Tracking is in your confirmation email.",
        ("h", "Damaged item"),
        "Hi {name}, sorry about that! Please send photos of the item and the box and we will make it right.",
    ])

    # D16 / D17 supplier agreements
    write_pdf("Supplier Agreement - Oakridge Textiles.pdf", "Supply Agreement: Oakridge Textiles and %s" % COMPANY, [
        "Term: 2026-01-01 to 2027-12-31.",
        ("h", "1. Orders"), "Minimum order quantity (MOQ) is 200 units per SKU per order.",
        ("h", "2. Lead time"), "Supplier ships within 35 days of a confirmed purchase order.",
        ("h", "3. Payment"), "Payment terms are net 45 from invoice date.",
        ("h", "4. Quality"), "Buyer may reject goods that fail the agreed specification within 14 days of receipt.",
    ])
    write_pdf("Supplier Agreement - Brightwood Furniture.pdf", "Supply Agreement: Brightwood Furniture and %s" % COMPANY, [
        "Term: 2025-07-01 to 2027-06-30.",
        ("h", "1. Lead time"), "Supplier ships within 60 days of a confirmed purchase order.",
        ("h", "2. Payment"), "Payment terms are net 30 from invoice date.",
        ("h", "3. Defects"), "If the defect rate in a quarter is over 2%, Supplier gives a credit equal to the cost of the defective units.",
    ])

    # D18
    write_xlsx("Supplier contacts.xlsx", [("Contacts", [
        ["Supplier", "Contact", "Role", "Email", "Reorder method"],
        ["Oakridge Textiles", "Dana Whitfield", "Account manager", "orders@oakridge-textiles.example", "Email PO"],
        ["Brightwood Furniture", "Luis Ortega", "Sales rep", "sales@brightwood-furniture.example", "Supplier portal"],
        ["Harbor Glassworks", "Mina Patel", "Owner", "hello@harbor-glass.example", "Email PO"],
    ])])

    # D19 / D20 sales
    for fname, q, vals in (("Q1 2026 Sales Summary.xlsx", "Q1 2026", (365800, 201200, 88100)),
                           ("Q2 2026 Sales Summary.xlsx", "Q2 2026", (412300, 186900, 97450))):
        write_xlsx(fname, [(q, [
            ["Category", "Revenue (USD)", "Orders"],
            ["Furniture", vals[0], int(vals[0] / 610)],
            ["Textiles", vals[1], int(vals[1] / 95)],
            ["Decor", vals[2], int(vals[2] / 48)],
        ])])

    # D21
    write_pdf("Product Care Guide.pdf", "Product Care Guide", [
        ("h", "Linen"), "Machine wash cold, gentle cycle. Tumble dry low or line dry.",
        ("h", "Oak tables"), "Wipe with a dry cloth. Apply furniture oil every 6 months. Keep out of direct sunlight.",
        ("h", "Glass decor"), "Hand wash only.",
    ])

    # D22 scanned lease
    write_scanned_pdf("Warehouse Lease (signed scan).pdf", [[
        "# COMMERCIAL LEASE AGREEMENT", "",
        "Landlord: Riverbend Industrial Properties LLC",
        "Tenant: %s" % COMPANY,
        "Premises: Unit 4, 1180 Millrace Road (warehouse, approx. 9,500 sq ft)", "",
        "# 1. Term",
        "The lease begins on July 1, 2024 and ends on June 30, 2027.", "",
        "# 2. Rent",
        "Monthly base rent is $8,200, due on the first day of each month.",
        "Late payments after the 5th incur a fee of $150.", "",
        "# 3. Security deposit",
        "Tenant has paid a security deposit of $16,400.", "",
        "Signed: ______________________   Date: 06/12/2024",
    ]])

    # D23 / D24 carrier conflict
    write_docx("Ops meeting notes 2026-05-12.docx", "Ops meeting notes - 2026-05-12", [
        "Attendees: Ops lead, CS lead, Owner",
        ("h", "Decisions"),
        ["Switch carrier from SwiftShip to ParcelPoint starting 2026-07-01 (better damage rates, lower cost).",
         "Keep SwiftShip for orders already in transit on that date."],
        ("h", "Action items"),
        ["Ops: sign ParcelPoint contract by 2026-06-01", "Marketing: update website shipping FAQ before July"],
    ])
    write_docx("Shipping FAQ (website copy).docx", "Shipping FAQ", [
        ("h", "Which carrier do you use?"), "We ship all orders with SwiftShip.",
        ("h", "When will my order ship?"), "Orders ship within 1-2 business days.",
        ("h", "Do you offer free shipping?"), "Yes, on qualifying US orders. See checkout for details.",
    ])

    # D25 - D29
    write_docx("Wholesale Inquiries Process.docx", "Wholesale Inquiries", [
        ["Minimum first order: $1,500.", "Wholesale discount: 20% off retail price.",
         "Send the wholesale application form. The owner approves new wholesale accounts."],
    ])
    write_xlsx("Promo Calendar 2026.xlsx", [("2026", [
        ["Promotion", "Start", "End", "Discount", "Scope"],
        ["Spring refresh", "2026-03-20", "2026-03-27", "15%", "Textiles"],
        ["Summer sale", "2026-07-10", "2026-07-17", "20%", "Selected decor"],
        ["Black Friday", "2026-11-27", "2026-11-30", "25%", "Sitewide"],
    ])])
    write_docx("Gift Card Terms.docx", "Gift Card Terms", [
        ["Gift cards do not expire.", "Gift cards cannot be redeemed for cash, except where the law requires it.",
         "Lost or stolen gift cards cannot be replaced."],
    ])
    write_pdf("Brand Voice Guide.pdf", "Brand Voice Guide", [
        "We sound like a friendly neighbor who knows a lot about homes: warm, clear, never pushy.",
        ["Use short sentences.", "Say 'we' and 'you'.", "No all-caps sales language."],
    ])
    write_docx("New CS Hire Onboarding Checklist.docx", "New CS Hire - Week 1", [
        ["Day 1: accounts for helpdesk and Shopify", "Day 2: read Return Policy and CS SOPs",
         "Day 3: shadow a senior agent", "Day 5: answer tickets with review"],
    ])

    # D30 policy index
    write_xlsx("Policy Index - current versions.xlsx", [("Index", [
        ["Topic", "Current file", "Effective", "Owner"],
        ["Returns", "Return Policy v3 final2.docx", "2026-04-15", "CS lead"],
        ["Shipping rates", "Shipping Rates 2026.xlsx", "2026-01-05", "Ops"],
        ["Employee handbook", "Employee Handbook 2026.pdf", "2026-01-01", "Owner"],
        ["Damaged items", "CS SOP - Damaged Items.docx", "2025-09-01", "CS lead"],
        ["Refund approvals", "CS SOP - Refund Escalation.docx", "2025-09-01", "CS lead"],
    ])])

    # D31 draft
    write_docx("Loyalty Program Proposal.docx", "Lantern Rewards", [
        "Customers earn 5% back in points on every order. 100 points = $1.",
        "Points expire after 12 months.",
        "Status: draft for owner review. Not approved.",
    ])

    print("%d files in %s" % (len(os.listdir(OUT)), OUT))


if __name__ == "__main__":
    build()
